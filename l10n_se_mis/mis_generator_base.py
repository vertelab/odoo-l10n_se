#!/usr/bin/env python3
"""
Gemensam basgenerator för MIS XML från Excel-taxonomi.
v3: parent_id, sum() på rader, svenska tecken.
"""
import os
from openpyxl import load_workbook

# ── Expanded BAS range overrides ──────────────────────────────────────────────
BAS_RANGE_EXPANSIONS = {
    '26xx': range(2600, 2670),
    '27xx': range(2700, 2800),
    '29xx': range(2900, 3000),
    '17xx': range(1700, 1800),
    '20xx': range(2000, 2100),
    '209x': range(2090, 2100),
    '208x': range(2080, 2090),
}


def expand_bas_range(bas_str):
    """Expandera BAS-konto range till lista med kontonummer."""
    if not bas_str:
        return []
    bas_str = bas_str.strip()
    if bas_str in BAS_RANGE_EXPANSIONS:
        return sorted(BAS_RANGE_EXPANSIONS[bas_str])
    accounts = []
    if '/' in bas_str:
        bas_str = bas_str.split('/')[-1]
    if '-' in bas_str:
        parts = bas_str.split('-')
        start, end = parts[0].strip(), parts[1].strip()
        if 'x' in start.lower() or 'x' in end.lower():
            for a in range(int(start.lower().replace('x', '0')),
                           int(end.lower().replace('x', '9')) + 1):
                accounts.append(a)
        else:
            for a in range(int(start), int(end) + 1):
                accounts.append(a)
    elif 'x' in bas_str.lower():
        base = int(bas_str.lower().replace('x', '0'))
        accounts.extend(base + i for i in range(10))
    else:
        try:
            accounts.append(int(bas_str))
        except ValueError:
            pass
    return sorted(set(accounts))


def compress_account_codes(accounts):
    """Kompakta kontonummer med %% wildcards för MIS Builder."""
    if not accounts:
        return ''
    from collections import defaultdict
    groups = defaultdict(list)
    for acc in accounts:
        groups[str(acc)[:2]].append(acc)
    result = []
    for prefix in sorted(groups.keys()):
        codes = sorted(groups[prefix])
        lo, hi = min(codes), max(codes)
        base = int(prefix + '00')
        if len(codes) == 100 and lo == base and hi == base + 99:
            result.append(f'{prefix}%%')
        else:
            result.extend(str(c) for c in codes)
    return ', '.join(result)


def format_bal_expression(accounts, negate=False):
    """Skapa MIS bal[]-uttryck med wildcard-kompression."""
    if not accounts:
        return ''
    expr = 'bal[' + compress_account_codes(accounts) + ']'
    return expr + ' * -1' if negate else expr


def extract_bas_accounts(sh, row_num):
    """Hitta BAS-kontoreferenser i en rad (kolumn 20+)."""
    accounts = []
    for c in range(20, min(sh.max_column + 1, 250) - 2):
        vc = str(sh.cell(row_num, c).value or '')
        vc1 = str(sh.cell(row_num, c + 1).value or '')
        if vc == 'BAS' and 'BAS-konto' in vc1:
            val = str(sh.cell(row_num, c + 2).value or '')
            if val:
                accounts.extend(expand_bas_range(val))
    return sorted(set(accounts))


def get_display_name(sh, row_num):
    """Hämta svenskt visningsnamn från kolumn E-H (5-8)."""
    for c in [8, 7, 6, 5]:
        val = sh.cell(row_num, c).value
        if val and str(val).strip():
            return str(val).strip()
    return ''


def get_hierarchy_level(sh, row_num):
    """Bestäm hierarkinivå (0-3) från kolumn E-H."""
    for c in [5, 6, 7, 8]:
        val = sh.cell(row_num, c).value
        if val and str(val).strip():
            return c - 5
    return None


def generate_styles_xml():
    """Standard 5-stilars XML."""
    return '''        <record id="report_style_k2_1" model="mis.report.style">
            <field name="name">Style for money K2</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
            <field eval="False" name="dp_inherit"/>
            <field name="dp">0</field>
            <field eval="False" name="indent_level_inherit"/>
            <field name="indent_level">2</field>
        </record>
        <record id="report_style_k2_2" model="mis.report.style">
            <field name="name">Style account K2</field>
            <field eval="False" name="indent_level_inherit"/>
            <field name="indent_level">4</field>
            <field eval="False" name="font_style_inherit"/>
            <field name="font_style">italic</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_3" model="mis.report.style">
            <field name="name">Style total K2</field>
            <field eval="False" name="background_color_inherit"/>
            <field name="background_color">#967C8B</field>
            <field eval="False" name="color_inherit"/>
            <field name="color">#FFFFFF</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="font_weight">bold</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_4" model="mis.report.style">
            <field name="name">Style bold K2</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="indent_level">0</field>
            <field name="font_weight">bold</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_5" model="mis.report.style">
            <field name="name">Style bold sum K2</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="indent_level">0</field>
            <field name="font_weight">bold</field>
            <field eval="False" name="font_size_inherit"/>
            <field name="font_size">medium</field>
            <field eval="False" name="prefix_inherit"/>
            <field name="suffix">SEK</field>
        </record>'''


def generate_report_xml(sheet_def, compact=False):
    """Generera XML för en rapport med parent_id och sum()-uttryck."""
    wb = load_workbook(sheet_def['excel_path'])
    sh = wb[sheet_def['sheet']]
    report_id = sheet_def['report_id']
    if compact:
        report_id += '_compact'
    report_name = sheet_def['report_name']
    if compact:
        report_name += ' (kompakt)'
    ec = sheet_def['elem_col']
    ac = sheet_def['abstract_col']
    sc = sheet_def['saldo_col']

    seq = [0]
    def nseq():
        seq[0] += 1
        return seq[0]

    # ── Läs alla rader ─────────────────────────────────────────────────────
    rows = []
    for r in range(2, sh.max_row + 1):
        elem = sh.cell(r, ec).value
        if not elem:
            continue
        elem = str(elem).strip()
        abstract = str(sh.cell(r, ac).value or '').lower() == 'true'
        saldo = str(sh.cell(r, sc).value or '')
        display = get_display_name(sh, r)
        level = get_hierarchy_level(sh, r)
        accounts = extract_bas_accounts(sh, r)
        is_sum = 'Summa' in display or elem.startswith('Summa')

        rows.append({
            'row': r, 'elem': elem, 'abstract': abstract,
            'display': display, 'level': level,
            'accounts': accounts, 'negate': saldo == 'credit',
            'is_sum': is_sum,
        })

    # ── Bygg parent_id-hierarki ───────────────────────────────────────────
    # parent = närmast föregående KPI med lägre level
    for i, rd in enumerate(rows):
        rd['parent_idx'] = None
        if rd['level'] is not None:
            for j in range(i - 1, -1, -1):
                p = rows[j]
                if p['level'] is not None and p['level'] < rd['level']:
                    rd['parent_idx'] = j
                    break

    # ── Generera sum()-uttryck för sum-KPI:er ─────────────────────────────
    # En sum-KPI på level N summerar KPI:er på level N+1 inom sitt sektion.
    # Section start = förra KPI:n med level < N.
    for i, rd in enumerate(rows):
        if not rd['is_sum']:
            continue
        if rd['accounts']:
            continue  # Har egna BAS-konton (leaf-KPI), behåll bal[]

        kpi_level = rd['level']
        if kpi_level is None:
            continue

        # Hitta section start: förra KPI:n med level < kpi_level
        section_start = 0
        for j in range(i - 1, -1, -1):
            p = rows[j]
            pl = p['level']
            if pl is not None and pl < kpi_level:
                section_start = j + 1
                break

        # Samla sum-barn: bara KPI:er på level == kpi_level + 1
        # Skippa level None (djupa löv) — de täcks av sub-sums på level N+1
        # Skippa nivåer djupare än kpi_level + 1 för att undvika dubbelräkning
        child_names = []
        for j in range(section_start, i):
            child = rows[j]
            cl = child['level']
            if child['abstract']:
                continue
            if cl is None:
                # Deep leaf without explicit level — skip, covered by sub-sums
                continue
            if cl == kpi_level + 1:
                child_names.append(child['elem'])

        if child_names:
            expr = "sum(" + ", ".join(f"'{n}'" for n in child_names) + ")"
            if rd['negate']:
                expr += ' * -1'
            rows[i]['sum_expr'] = expr
            rows[i]['accounts'] = []  # Rensa bal[] (använd sum())

    # ─── XML-utdata ──────────────────────────────────────────────────────
    lines = []
    lines.append(f'''        <record id="{report_id}" model="mis.report">
            <field name="name">{report_name}</field>
        </record>''')

    for rd in rows:
        s = nseq()
        kpi_id = f"{report_id}_{rd['elem']}"

        # Stil:
        # - level 1 abstract headers (huvudrubriker) → style_3 (grå bakgrund)
        # - Other abstract headers → style_4 (bold)
        # - Sum rows (non-abstract) → style_5 (bold sum, medium)
        # - Leaf KPIs → style_1 (money, indent 2)
        if rd['abstract'] and rd['level'] == 1:
            style = 'report_style_k2_3'
        elif rd['abstract']:
            style = 'report_style_k2_4'
        elif rd['is_sum']:
            style = 'report_style_k2_5'
        else:
            style = 'report_style_k2_1'

        # Typ
        type_field = ''
        budgetable = 'True'
        if rd['abstract'] and not rd['is_sum']:
            type_field = '\n            <field name="type">str</field>'
            budgetable = 'False'
        elif rd['is_sum']:
            type_field = '\n            <field name="type">num</field>'
            budgetable = 'False'

        # auto_expand_accounts (bara leaf-KPI:er i full version)
        auto = ''
        auto_style = ''
        if not compact and not rd['abstract'] and not rd['is_sum']:
            auto = '\n            <field name="auto_expand_accounts">True</field>'
            auto_style = '\n            <field name="auto_expand_accounts_style_id" ref="report_style_k2_2"/>'

        # Parent
        parent_field = ''
        if rd['parent_idx'] is not None:
            p = rows[rd['parent_idx']]
            parent_id = f"{report_id}_{p['elem']}"
            parent_field = f'\n            <field name="parent_id" ref="{parent_id}"/>'

        display = rd['display'] or rd['elem']

        lines.append(f'''        <record id="{kpi_id}" model="mis.report.kpi">
            <field name="report_id" ref="{report_id}"/>
            <field name="name">{rd['elem']}</field>
            <field name="description">{display}</field>
            <field name="style_id" ref="{style}"/>
            <field name="sequence">{s}</field>{type_field}
            <field name="budgetable">{budgetable}</field>{auto}{auto_style}{parent_field}
        </record>''')

        # Expression
        if rd.get('sum_expr'):
            eid = f"kpi_{kpi_id}"
            lines.append(f'''        <record id="{eid}" model="mis.report.kpi.expression">
            <field name="kpi_id" ref="{kpi_id}"/>
            <field name="name">{rd['sum_expr']}</field>
        </record>''')
        elif rd['accounts'] and not rd['is_sum']:
            bal = format_bal_expression(rd['accounts'], rd['negate'])
            eid = f"kpi_{kpi_id}"
            lines.append(f'''        <record id="{eid}" model="mis.report.kpi.expression">
            <field name="kpi_id" ref="{kpi_id}"/>
            <field name="name">{bal}</field>
        </record>''')

    wb.close()
    return '\n'.join(lines)


def write_xml(filename, output_dir, styles_xml, report_xml):
    full = f'''<?xml version="1.0" encoding="utf-8"?>
<odoo>
    <data noupdate="1">
{styles_xml}

{report_xml}
    </data>
</odoo>'''
    path = os.path.join(output_dir, filename)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(full)
    return os.path.getsize(path)


def generate_all(sheets, output_dir):
    print("Generating MIS XML files from Excel taxonomy...")
    styles = generate_styles_xml()
    for sd in sheets:
        print(f"  Processing sheet: {sd['sheet']}...")
        r = generate_report_xml(sd, compact=False)
        s = write_xml(sd['filename'], output_dir, styles, r)
        print(f"    → Full: {sd['filename']} ({s} bytes)")
        r2 = generate_report_xml(sd, compact=True)
        s2 = write_xml(sd['filename_compact'], output_dir, styles, r2)
        print(f"    → Compact: {sd['filename_compact']} ({s2} bytes)")
    print("Done!")

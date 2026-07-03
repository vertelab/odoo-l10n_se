#!/usr/bin/env python3
"""
Generator: Läser K2-taxonomi Excel och genererar MIS XML datafiler.
Kör: python3 bin/generate_mis_xml.py
"""
import os, re
from openpyxl import load_workbook

EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                          'k2-filial-arsredovisning-2024-09-12_rev20250312_sv.xlsx')
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data')
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ── Sheet definitions ──────────────────────────────────────────────────────────
SHEETS = [
    {
        'sheet': 'Balansräkning',
        'report_id': 'report_br_filial_k2',
        'report_name': 'Balansräkning K2',
        'filename': 'mis_balansrakning_k2.xml',
        'filename_compact': 'mis_balansrakning_k2_compact.xml',
        'elem_col': 10,
        'abstract_col': 13,
        'saldo_col': 15,
    },
    {
        'sheet': 'Förkortad balansräkning',
        'report_id': 'report_br_filial_forkortad_k2',
        'report_name': 'Förkortad balansräkning K2',
        'filename': 'mis_balansrakning_forkortad_k2.xml',
        'filename_compact': 'mis_balansrakning_forkortad_k2_compact.xml',
        'elem_col': 10,
        'abstract_col': 13,
        'saldo_col': 15,
    },
    {
        'sheet': 'Kostnadsslagsindelad resultatr',
        'report_id': 'report_rr_filial_kostnadsslag_k2',
        'report_name': 'Resultaträkning kostnadsslagsindelad K2',
        'filename': 'mis_resultatrakning_kostnadsslag_k2.xml',
        'filename_compact': 'mis_resultatrakning_kostnadsslag_k2_compact.xml',
        'elem_col': 9,
        'abstract_col': 12,
        'saldo_col': 14,
    },
    {
        'sheet': 'Förkortad kostnadsslagsindelad',
        'report_id': 'report_rr_filial_kostnadsslag_forkortad_k2',
        'report_name': 'Förkortad resultaträkning kostnadsslagsindelad K2',
        'filename': 'mis_resultatrakning_kostnadsslag_forkortad_k2.xml',
        'filename_compact': 'mis_resultatrakning_kostnadsslag_forkortad_k2_compact.xml',
        'elem_col': 8,
        'abstract_col': 11,
        'saldo_col': 13,
    },
    {
        'sheet': 'Kassaflödesanalys indirekt met',
        'report_id': 'report_kf_filial_indirekt_k2',
        'report_name': 'Kassaflödesanalys indirekt metod K2',
        'filename': 'mis_kassaflodesanalys_k2.xml',
        'filename_compact': 'mis_kassaflodesanalys_k2_compact.xml',
        'elem_col': 9,
        'abstract_col': 12,
        'saldo_col': 14,
    },
]

# ── Helpers ────────────────────────────────────────────────────────────────────

def expand_bas_range(bas_str):
    """Expandera BAS-konto range som '191x-198x' eller '1331-1334' till lista."""
    if not bas_str:
        return []
    bas_str = bas_str.strip()
    accounts = []
    if '/' in bas_str:
        bas_str = bas_str.split('/')[-1]
    if '-' in bas_str:
        parts = bas_str.split('-')
        start = parts[0].strip()
        end = parts[1].strip()
        if 'x' in start.lower() or 'x' in end.lower():
            start_num = int(start.lower().replace('x', '0'))
            end_num_str = end.lower().replace('x', '9')
            end_num = int(end_num_str)
            for a in range(start_num, end_num + 1):
                accounts.append(a)
        else:
            start_num = int(start)
            end_num = int(end)
            for a in range(start_num, end_num + 1):
                accounts.append(a)
    elif 'x' in bas_str.lower():
        base = int(bas_str.lower().replace('x', '0'))
        for i in range(10):
            accounts.append(base + i)
    else:
        try:
            accounts.append(int(bas_str))
        except ValueError:
            pass
    return sorted(set(accounts))


def format_bal_expression(accounts, negate=False):
    """Skapa MIS bal[]-uttryck från kontonummer."""
    if not accounts:
        return ''
    expr = 'bal[' + ', '.join(str(a) for a in accounts) + ']'
    if negate:
        expr += ' * -1'
    return expr


def extract_bas_accounts(sh, row_num):
    """Hitta BAS-kontoreferenser i en rad.
    
    BAS-referenser finns som grupper av 3 kolumner:
      Col C: 'BAS' (Utgivare)
      Col C+1: 'BAS-konto' (Namn)
      Col C+2: kontorangen (Nummer), t.ex. '191x-198x'
    """
    accounts = []
    max_col = min(sh.max_column + 1, 250)
    for c in range(20, max_col - 2):
        val_c = str(sh.cell(row_num, c).value or '')
        val_c1 = str(sh.cell(row_num, c + 1).value or '')
        if val_c == 'BAS' and 'BAS-konto' in val_c1:
            num_val = str(sh.cell(row_num, c + 2).value or '')
            if num_val:
                accounts.extend(expand_bas_range(num_val))
    return sorted(set(accounts))


def get_display_name(sh, row_num):
    """Hämta det svenska visningsnamnet för en KPI."""
    for c in [8, 7, 6, 5]:
        val = sh.cell(row_num, c).value
        if val and str(val).strip():
            return str(val).strip()
    return ''


# ── Main generator ─────────────────────────────────────────────────────────────

def generate_styles_xml():
    return '''        <record id="report_style_k2_1" model="mis.report.style">
            <field name="name">Style for money K2</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
            <field eval="False" name="dp_inherit"/>
            <field name="dp">0</field>
            <field eval="False" name="indent_level_inherit"/>
            <field name="indent_level">2</field>
        </record>
        <record id="report_style_k2_2" model="mis.report.style">
            <field name="name">Style account K2</field>
            <field eval="False" name="indent_level_inherit"/>
            <field name="indent_level">4</field>
            <field eval="False" name="font_style_inherit"/>
            <field name="font_style">italic</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_3" model="mis.report.style">
            <field name="name">Style total K2</field>
            <field eval="False" name="background_color_inherit"/>
            <field name="background_color">#967C8B</field>
            <field eval="False" name="color_inherit"/>
            <field name="color">#FFFFFF</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="font_weight">bold</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_4" model="mis.report.style">
            <field name="name">Style bold K2</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="indent_level">0</field>
            <field name="font_weight">bold</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>
        <record id="report_style_k2_5" model="mis.report.style">
            <field name="name">Style bold sum K2</field>
            <field eval="False" name="font_weight_inherit"/>
            <field name="indent_level">0</field>
            <field name="font_weight">bold</field>
            <field eval="False" name="font_size_inherit"/>
            <field name="font_size">medium</field>
            <field eval="False" name="prefix_inherit"/>
            <field eval="False" name="suffix_inherit"/>
            <field name="suffix">SEK</field>
        </record>'''


def generate_report_xml(sheet_def, compact=False):
    """Generera XML för en rapport.
    compact=True: utelämnar auto_expand_accounts (inga kontokolumner)."""
    wb = load_workbook(EXCEL_PATH)
    sh = wb[sheet_def['sheet']]
    report_id = sheet_def['report_id']
    if compact:
        report_id += '_compact'
    report_name = sheet_def['report_name']
    if compact:
        report_name += ' (kompakt)'
    elem_col = sheet_def['elem_col']
    abstract_col = sheet_def['abstract_col']
    saldo_col = sheet_def['saldo_col']

    lines = []
    seq_counter = [0]

    def next_seq():
        seq_counter[0] += 1
        return seq_counter[0]

    rows_data = []
    for r in range(2, sh.max_row + 1):
        elem = sh.cell(r, elem_col).value
        if elem and str(elem).strip():
            abstract = str(sh.cell(r, abstract_col).value or '').lower() == 'true'
            saldo = str(sh.cell(r, saldo_col).value or '')
            display = get_display_name(sh, r)
            rows_data.append({
                'row': r,
                'elem': str(elem).strip(),
                'abstract': abstract,
                'saldo': saldo,
                'display': display,
            })

    # ─── Report record ─────────────────────────────────────────────────────
    lines.append(f'''        <record id="{report_id}" model="mis.report">
            <field name="name">{report_name}</field>
        </record>''')

    for rd in rows_data:
        seq = next_seq()
        elem = rd['elem']
        is_abstract = rd['abstract']
        saldo = rd['saldo']
        display = rd['display'] or elem
        kpi_id = f"{report_id}_{elem}"

        is_sum = 'Summa' in display or elem.startswith('Summa')

        if is_abstract and is_sum:
            style = 'report_style_k2_5'
        elif is_abstract:
            style = 'report_style_k2_4'
        else:
            style = 'report_style_k2_1'

        type_field = ''
        budgetable = 'True'
        if is_abstract and not is_sum:
            type_field = '\n            <field name="type">str</field>'
            budgetable = 'False'
        elif is_sum:
            type_field = '\n            <field name="type">num</field>'
            budgetable = 'False'

        # auto_expand_accounts only in full version, not compact
        auto_expand = ''
        auto_style = ''
        if not compact and not is_abstract and not is_sum:
            auto_expand = '\n            <field name="auto_expand_accounts">True</field>'
            auto_style = '\n            <field name="auto_expand_accounts_style_id" ref="report_style_k2_2"/>'

        lines.append(f'''        <record id="{kpi_id}" model="mis.report.kpi">
            <field name="report_id" ref="{report_id}"/>
            <field name="name">{elem}</field>
            <field name="description">{display}</field>
            <field name="style_id" ref="{style}"/>
            <field name="sequence">{seq}</field>{type_field}
            <field name="budgetable">{budgetable}</field>{auto_expand}{auto_style}
        </record>''')

        accounts = extract_bas_accounts(sh, rd['row'])
        if accounts:
            negate = saldo == 'credit'
            bal_expr = format_bal_expression(accounts, negate)
            expr_id = f"kpi_{kpi_id}"
            lines.append(f'''        <record id="{expr_id}" model="mis.report.kpi.expression">
            <field name="kpi_id" ref="{kpi_id}"/>
            <field name="name">{bal_expr}</field>
        </record>''')

    wb.close()
    return '\n'.join(lines)


def write_xml(filename, styles_xml, report_xml):
    """Write one XML file with styles + report data."""
    full_xml = f'''<?xml version="1.0" encoding="utf-8"?>
<odoo>
    <data noupdate="1">
{styles_xml}

{report_xml}
    </data>
</odoo>'''
    filepath = os.path.join(OUTPUT_DIR, filename)
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(full_xml)
    return os.path.getsize(filepath)


def generate_module_xml():
    print("Generating MIS XML files from Excel taxonomy...")
    styles_xml = generate_styles_xml()
    for sd in SHEETS:
        print(f"  Processing sheet: {sd['sheet']}...")
        # Full version (with auto_expand_accounts)
        report_xml = generate_report_xml(sd, compact=False)
        size = write_xml(sd['filename'], styles_xml, report_xml)
        print(f"    → Full: {sd['filename']} ({size} bytes)")
        # Compact version (no auto_expand_accounts)
        report_xml_compact = generate_report_xml(sd, compact=True)
        size_c = write_xml(sd['filename_compact'], styles_xml, report_xml_compact)
        print(f"    → Compact: {sd['filename_compact']} ({size_c} bytes)")
    print("\nDone!")


if __name__ == '__main__':
    generate_module_xml()
